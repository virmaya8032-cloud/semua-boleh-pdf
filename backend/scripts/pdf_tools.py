"""Local PDF processing. JSON is received on stdin; documents never leave the server."""
import base64
import json
import math
import os
import sys
import zipfile
from decimal import Decimal, InvalidOperation
import fitz


def number(options, key, default, minimum, maximum):
    value = options.get(key)
    try:
        result = float(default if value in (None, '') else value)
    except (ValueError, TypeError):
        raise ValueError(f'Nilai {key} mesti nombor.')
    if not math.isfinite(result) or not minimum <= result <= maximum:
        raise ValueError(f'Nilai {key} mesti antara {minimum} dan {maximum}.')
    return result


def text(options, key, default='', required=False):
    result = str(options.get(key, default) or '').strip()
    if required and not result:
        raise ValueError(f'Sila isi {key.replace("_", " ")}.')
    if len(result) > 30000:
        raise ValueError('Teks terlalu panjang.')
    return result


def opened(filename):
    doc = fitz.open(filename)
    if not doc.is_pdf or doc.needs_pass or not len(doc):
        doc.close()
        raise ValueError('Perlu PDF yang sah, mempunyai halaman dan tidak berkunci.')
    return doc


def save(doc, filename, fresh=False):
    if fresh:
        # Copy only live pages into a new document: no hidden attachments or stale objects.
        clean = fitz.open()
        clean.insert_pdf(doc, links=False, annots=False, widgets=False)
        clean.save(filename, garbage=4, deflate=True, clean=True)
        clean.close()
    else:
        doc.save(filename, garbage=4, deflate=True)


def rectangle(page, item):
    w, h = page.rect.width, page.rect.height
    x = number(item, 'x', 0, 0, 100)
    y = number(item, 'y', 20, 0, 100)
    width = number(item, 'lebar', 100 - x, 0.1, 100)
    height = number(item, 'tinggi', 5, 0.1, 100)
    return fitz.Rect(w*x/100, h*y/100, w*min(100,x+width)/100, h*min(100,y+height)/100) * page.derotation_matrix


def page_for(doc, item):
    index = int(number(item, 'halaman', 1, 1, len(doc))) - 1
    return doc[index]


def items_from(options, key='anotasi'):
    try:
        items = json.loads(options.get(key, '[]'))
    except (ValueError, TypeError):
        raise ValueError('Senarai anotasi tidak sah.')
    if not isinstance(items, list) or len(items) > 200:
        raise ValueError('Maksimum 200 anotasi dibenarkan.')
    if any(not isinstance(item, dict) for item in items):
        raise ValueError('Anotasi tidak sah.')
    return items


def page_size(options):
    sizes = {'a4': (595.276, 841.89), 'a3': (841.89, 1190.551), 'a5': (419.528, 595.276), 'letter': (612, 792), 'legal': (612, 1008)}
    key = options.get('saiz_kertas', 'a4')
    if key not in sizes:
        raise ValueError('Saiz kertas tidak sah.')
    w, h = sizes[key]
    return (h, w) if options.get('orientasi') == 'landskap' else (w, h)


def paragraphs(doc, title, lines):
    """Wrap long user text and add pages rather than clip it."""
    page = doc.new_page(width=595.276, height=841.89)
    y = 50
    for value, size, bold in [(title, 20, True)] + lines:
        font = fitz.Font('hebo' if bold else 'helv')
        for raw in str(value).splitlines() or ['']:
            words = raw.split()
            line = ''
            wrapped = []
            for word in words:
                # Long unbroken values are split to keep the output within the page.
                for piece in [word[i:i+65] for i in range(0, len(word), 65)]:
                    candidate = (line+' '+piece).strip()
                    if font.text_length(candidate, fontsize=size) > 495 and line:
                        wrapped.append(line)
                        line = piece
                    else:
                        line = candidate
            wrapped.append(line)
            for row in wrapped:
                if y + size*1.5 > 790:
                    page = doc.new_page(width=595.276, height=841.89)
                    y = 50
                page.insert_text((50,y), row, fontsize=size, fontname='hebo' if bold else 'helv')
                y += size*1.5
        y += 7
    return doc


def process(op, paths, opts, output):
    if op in ('invois', 'permohonan-kerja'):
        doc = fitz.open()
        if op == 'invois':
            supplier = text(opts, 'penjual', required=True)
            customer = text(opts, 'pelanggan', required=True)
            raw_items = text(opts, 'item', required=True).splitlines()
            if len(raw_items) > 100:
                raise ValueError('Maksimum 100 item invois.')
            lines = [(supplier, 12, True), (f"Pelanggan: {customer}", 11, False), (f"No. invois: {text(opts,'nombor','INV-001')}", 11, False), (f"Tarikh: {text(opts,'tarikh')}", 11, False)]
            total = Decimal('0')
            for i, line in enumerate(raw_items, 1):
                parts = line.rsplit('|', 2)
                if len(parts) != 3:
                    raise ValueError('Setiap baris item mesti: Perkara | Kuantiti | Harga seunit')
                name, qty, price = parts
                try:
                    qty, price = Decimal(qty.strip()), Decimal(price.strip())
                    if not qty.is_finite() or not price.is_finite() or qty <= 0 or price < 0 or qty > 1000000 or price > 100000000:
                        raise InvalidOperation()
                except InvalidOperation:
                    raise ValueError('Kuantiti dan harga item tidak sah.')
                amount = (qty*price).quantize(Decimal('0.01'))
                total += amount
                lines.append((f'{i}. {name.strip()} — {qty} x RM{price:.2f} = RM{amount:.2f}', 11, False))
            lines += [(f'JUMLAH: RM{total:.2f}', 16, True), (text(opts,'nota'), 11, False)]
            paragraphs(doc, 'INVOIS', lines)
        else:
            name = text(opts, 'nama', required=True)
            role = text(opts, 'jawatan', required=True)
            company = text(opts, 'syarikat', required=True)
            lines = [(name, 13, True), (text(opts,'alamat'), 11, False), (f"E-mel: {text(opts,'emel')}   Telefon: {text(opts,'telefon')}", 11, False), (text(opts,'tarikh'), 11, False), (f'Kepada: {company}', 12, True), (f'PERMOHONAN JAWATAN {role.upper()}', 13, True), (text(opts,'surat',required=True), 11, False), ('Yang benar,', 11, False), (name, 12, True)]
            paragraphs(doc, 'SURAT PERMOHONAN KERJA', lines)
        save(doc, output)
        doc.close()
        return

    if op == 'banding':
        import difflib
        first,second = opened(paths[0]),opened(paths[1])
        try:
            def lines(document):
                result = []
                for i,page in enumerate(document):
                    result.append(f'--- Halaman {i+1} ---\n')
                    result += [line+'\n' for line in page.get_text(sort=True).splitlines()]
                return result
            changes = ''.join(difflib.unified_diff(lines(first),lines(second),fromfile='Fail 1',tofile='Fail 2'))
            with open(output,'w',encoding='utf-8') as target:
                target.write('LAPORAN PERBANDINGAN TEKS PDF\n\n'+(changes or 'Tiada perbezaan teks ditemui.\n'))
        finally:
            first.close();second.close()
        return
    doc = opened(paths[0]) if paths else fitz.open()
    try:
        if op in ('pdf-ke-word', 'pdf-ke-excel', 'pdf-ke-powerpoint'):
            if op == 'pdf-ke-word':
                from docx import Document
                result = Document()
                for index,page in enumerate(doc):
                    if index: result.add_page_break()
                    blocks = page.get_text('blocks',sort=True)
                    for block in blocks:
                        if block[6] == 0 and block[4].strip(): result.add_paragraph(block[4].strip())
                if not any(p.text.strip() for p in result.paragraphs):
                    raise ValueError('Tiada teks boleh diekstrak. Jalankan OCR pada PDF dahulu.')
                result.save(output)
            elif op == 'pdf-ke-excel':
                from openpyxl import Workbook
                result = Workbook()
                result.remove(result.active)
                has_content = False
                for index,page in enumerate(doc):
                    tables = page.find_tables().tables
                    if tables:
                        for ti,table in enumerate(tables):
                            sheet = result.create_sheet(f'H{index+1}-J{ti+1}')
                            for row in table.extract():
                                # Force string cell types: PDF text starting with '=' is not a spreadsheet formula.
                                sheet.append(row)
                                for cell in sheet[sheet.max_row]:
                                    if isinstance(cell.value,str): cell.data_type = 's'
                            has_content = True
                    else:
                        content = page.get_text(sort=True).strip()
                        if content:
                            sheet = result.create_sheet(f'Halaman {index+1}')
                            sheet.append(['Teks PDF (tiada jadual dikenal pasti)'])
                            for line in content.splitlines():
                                sheet.append([line])
                                sheet.cell(sheet.max_row,1).data_type = 's'
                            has_content = True
                if not has_content:
                    raise ValueError('Tiada teks atau jadual. Jalankan OCR pada PDF dahulu.')
                result.save(output)
            else:
                from pptx import Presentation
                from pptx.util import Pt
                from io import BytesIO
                result = Presentation()
                result.slide_width = Pt(doc[0].rect.width)
                result.slide_height = Pt(doc[0].rect.height)
                for page in doc:
                    slide = result.slides.add_slide(result.slide_layouts[6])
                    pix = page.get_pixmap(dpi=150,alpha=False)
                    scale = min(result.slide_width/page.rect.width,result.slide_height/page.rect.height)
                    width,height = int(page.rect.width*scale),int(page.rect.height*scale)
                    slide.shapes.add_picture(BytesIO(pix.tobytes('png')),int((result.slide_width-width)/2),int((result.slide_height-height)/2),width=width,height=height)
                result.save(output)
        elif op == 'sensor':
            doc.bake()  # Render annotations/widgets before redacting their visible content.
            items = items_from(opts)
            if not items:
                items = [opts]
            for item in items:
                page = page_for(doc, item)
                page.add_redact_annot(rectangle(page, item), fill=(0,0,0), cross_out=False)
            for page in doc:
                # Remove text, overlapping vectors, and pixels inside image redactions.
                page.apply_redactions(images=2, graphics=2, text=0)
            save(doc, output, fresh=True)
        elif op == 'edit':
            items = items_from(opts)
            if not items:
                raise ValueError('Tambah sekurang-kurangnya satu teks, gambar atau anotasi.')
            for item in items:
                page = page_for(doc, item)
                kind = item.get('jenis')
                rect = rectangle(page, item)
                if kind == 'text':
                    value = text(item,'teks',required=True)
                    size = number(item,'saiz',14,6,72)
                    # Built-in CJK font handles Unicode on the PDF page.
                    remaining = page.insert_textbox(rect, value, fontsize=size, rotate=page.rotation, fontname='china-s', color=(0,0,0), overlay=True)
                    if remaining < 0:
                        raise ValueError('Kotak teks terlalu kecil. Besarkan kotak atau kecilkan saiz teks.')
                elif kind == 'highlight':
                    annot = page.add_highlight_annot(rect)
                    annot.set_opacity(0.35)
                    annot.update()
                elif kind == 'box':
                    page.draw_rect(rect, color=(0.1,0.2,0.8), width=1.5, overlay=True)
                elif kind == 'pen':
                    points = item.get('points', [])
                    if not isinstance(points,list) or len(points) < 2 or len(points) > 5000:
                        raise ValueError('Lukisan tidak sah.')
                    w,h = page.rect.width,page.rect.height
                    transformed = [fitz.Point(number(p,'x',0,0,100)*w/100,number(p,'y',0,0,100)*h/100)*page.derotation_matrix for p in points]
                    shape = page.new_shape()
                    shape.draw_polyline(transformed)
                    shape.finish(color=(0,0,0), width=2)
                    shape.commit(overlay=True)
                elif kind == 'image':
                    data = text(item,'data',required=True)
                    if not data.startswith(('data:image/png;base64,','data:image/jpeg;base64,')) or len(data)>4000000:
                        raise ValueError('Gambar anotasi mesti PNG/JPG, maksimum 3 MB.')
                    page.insert_image(rect, stream=base64.b64decode(data.split(',',1)[1], validate=True), keep_proportion=True, overlay=True)
                else:
                    raise ValueError('Jenis anotasi tidak disokong.')
            doc.bake()
            save(doc,output)
        elif op == 'ekstrak-gambar':
            count = 0
            seen = set()
            with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as archive:
                for page in doc:
                    for image in page.get_images(full=True):
                        xref, mask = image[0],image[1]
                        if xref in seen:
                            continue
                        seen.add(xref)
                        pix = fitz.Pixmap(doc,xref)
                        if mask:
                            pix = fitz.Pixmap(pix,fitz.Pixmap(doc,mask))
                        if pix.colorspace and pix.colorspace.n > 3:
                            pix = fitz.Pixmap(fitz.csRGB,pix)
                        count += 1
                        archive.writestr(f'gambar-{count:03d}.png',pix.tobytes('png'))
                if not count:
                    raise ValueError('Tiada gambar terbenam ditemui dalam PDF.')
        elif op in ('overlay','nup','saiz-halaman'):
            doc.bake()
            out = fitz.open()
            try:
                if op == 'overlay':
                    overlay = opened(paths[1])
                    overlay.bake()
                    try:
                        mode = opts.get('mod','ulang')
                        if mode == 'sepadan' and len(overlay) != len(doc):
                            raise ValueError('Bilangan halaman kedua-dua PDF mesti sama.')
                        for i, source in enumerate(doc):
                            target = out.new_page(width=source.rect.width,height=source.rect.height)
                            if source.get_contents():
                                target.show_pdf_page(target.rect,doc,i)
                            idx = i % len(overlay)
                            if overlay[idx].get_contents():
                                target.show_pdf_page(target.rect,overlay,idx,overlay=True)
                    finally:
                        overlay.close()
                else:
                    w,h = page_size(opts)
                    margin = number(opts,'margin_pt',12,0,100)
                    per_page = int(number(opts,'bilangan',2,1,16)) if op == 'nup' else 1
                    columns = math.ceil(math.sqrt(per_page))
                    rows = math.ceil(per_page/columns)
                    cw,ch = (w-2*margin)/columns,(h-2*margin)/rows
                    for i,page in enumerate(doc):
                        if i % per_page == 0:
                            target = out.new_page(width=w,height=h)
                        cell = i % per_page
                        x,y = margin+(cell%columns)*cw,margin+(cell//columns)*ch
                        rect = fitz.Rect(x+3,y+3,x+cw-3,y+ch-3)
                        if page.get_contents():
                            target.show_pdf_page(rect,doc,i,keep_proportion=True)
                save(out,output)
            finally:
                out.close()
        elif op in ('edit-metadata','buang-metadata'):
            if op == 'buang-metadata':
                doc.set_metadata({key: "" for key in ("title", "author", "subject", "keywords", "creator", "producer", "creationDate", "modDate", "trapped")})
                doc.del_xml_metadata()
            else:
                metadata = dict(doc.metadata)
                for key in ('title','author','subject','keywords'):
                    metadata[key] = text(opts,key)
                doc.set_metadata(metadata)
            save(doc,output)
        elif op == 'flatten':
            doc.bake()
            save(doc,output)
        elif op == 'cipta-borang':
            if not len(doc):
                doc.new_page(width=595.276,height=841.89)
            fields = items_from(opts,'medan')
            if not fields:
                raise ValueError('Tambah sekurang-kurangnya satu medan borang.')
            names = {widget.field_name for page in doc for widget in page.widgets()}
            for item in fields:
                page = page_for(doc,item)
                name = text(item,'nama',required=True)
                if name in names:
                    raise ValueError(f'Nama medan pendua: {name}')
                names.add(name)
                widget = fitz.Widget()
                widget.field_name = name
                widget.field_label = name
                widget.field_type = fitz.PDF_WIDGET_TYPE_CHECKBOX if item.get('jenis') == 'checkbox' else fitz.PDF_WIDGET_TYPE_TEXT
                widget.rect = rectangle(page,item)
                widget.text_font = 'Helv'
                widget.text_fontsize = 11
                widget.border_color = (0.45,0.45,0.45)
                widget.border_width = 1
                widget.fill_color = (0.95,0.97,1)
                page.add_widget(widget)
            save(doc,output)
        else:
            raise ValueError('Operasi PDF tidak dikenali.')
    finally:
        doc.close()


if __name__ == '__main__':
    output = None
    try:
        request = json.load(sys.stdin)
        output = request['output']
        process(request['op'],request['paths'],request.get('opts',{}),output)
        print(json.dumps({'ok':True}))
    except Exception as error:
        if output and os.path.isfile(output):
            os.unlink(output)
        print(json.dumps({'error':str(error)},ensure_ascii=False))
        sys.exit(1)
